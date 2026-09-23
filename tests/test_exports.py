"""Focused tests for literal content, provenance and export readability structure."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import sys
import unittest
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from exports import render_docx, render_markdown, _transcript_paragraphs
from docx import Document


FIXTURE = {
    "id": "synthetic-test-only",
    "title": "Synthetic example meeting",
    "filename": "synthetic-example-recording.wav",
    "created_at": "2026-09-22T10:30:00+00:00",
    "duration": 195.5,
    "language": "en",
    "speakers": [{"id": "SPEAKER_00", "name": "Example Alex"},
                 {"id": "SPEAKER_01", "name": "Example Taylor"}],
    "segments": [
        {"id": 0, "start": 0, "end": 34.1, "speaker": "SPEAKER_00",
         "text": "This is a synthetic test recording. Let us try the new meeting notes workflow with a short example first."},
        {"id": 1, "start": 34.1, "end": 76.0, "speaker": "SPEAKER_01",
         "text": "Agreed. We will use the small test recording for the first review. I will check the speaker labels tomorrow."},
        {"id": 2, "start": 76.0, "end": 123.0, "speaker": "SPEAKER_00",
         "text": "I will listen to the audio and correct any missing words. We should verify each name before exporting the meeting notes."},
        {"id": 3, "start": 123.0, "end": 169.0, "speaker": "SPEAKER_01",
         "text": "How should we handle the section where both people speak at the same time? The labels might need a manual correction."},
        {"id": 4, "start": 169.0, "end": 195.5, "speaker": "SPEAKER_00",
         "text": "That is still open. We can inspect it during the review and decide then."},
    ],
    "notes": {
        "overview": "This synthetic example describes a first review of the meeting notes workflow. The speakers agreed to use a short recording and to check the transcript and speaker labels before exporting.",
        "decisions": [{"text": "Use the small test recording for the first review.", "segment_ids": [1]}],
        "actions": [{"text": "Check the speaker labels.", "owner": "SPEAKER_01", "due": "Tomorrow as stated in the recording", "segment_ids": [1]},
                    {"text": "Listen to the audio and correct missing words.", "owner": "SPEAKER_00", "due": "", "segment_ids": [2]}],
        "open_questions": [{"text": "How should overlapping speech be handled?", "segment_ids": [3, 4]}],
    },
    "notes_stale": False,
    "warnings": [],
}


class ExportTests(unittest.TestCase):
    def test_markdown_literal_source_cannot_create_links_html_or_headings(self):
        job = deepcopy(FIXTURE)
        job["segments"][0]["text"] = "# injected\n<script>alert(1)</script> [open](https://bad.test) **words**"
        output = render_markdown(job)
        self.assertIn("\\# injected", output)
        self.assertIn("&lt;script&gt;", output)
        self.assertIn("\\[open\\]\\(https://bad\\.test\\)", output)
        self.assertNotIn("<script>", output)

    def test_notes_sources_use_segment_ids_and_current_names(self):
        job = deepcopy(FIXTURE)
        job["speakers"][1]["name"] = "New name"
        output = render_markdown(job)
        self.assertIn("Owner: New name", output)
        self.assertIn("Source: 00:34", output)
        self.assertIn("02:03, 02:49", output)
        self.assertIn("due: Not specified", output)

    def test_missing_sources_are_visible_without_fabricated_timestamps(self):
        job = deepcopy(FIXTURE)
        job["notes"]["decisions"][0]["segment_ids"] = [999]
        self.assertIn("Source reference unavailable", render_markdown(job))

    def test_stale_warning_and_import_metadata_are_truthful(self):
        job = deepcopy(FIXTURE)
        job["notes_stale"] = True
        output = render_markdown(job)
        self.assertIn("Notes need updating", output)
        self.assertIn("**Imported:**", output)
        self.assertNotIn("Meeting date", output)
        self.assertIn("UTC", output)

    def test_docx_preserves_literals_and_has_separate_transcript_and_headings(self):
        job = deepcopy(FIXTURE)
        job["notes_stale"] = True
        job["segments"][0]["text"] = "Português y español & <literal> \x00\x01\ud800 content"
        data = render_docx(job)
        document = Document(BytesIO(data))
        content = "\n".join(p.text for p in document.paragraphs)
        self.assertIn("Português y español & <literal>", content)
        self.assertIn("Notes need updating", content)
        self.assertIn("00:34  Example Taylor", content)
        transcript = next(p for p in document.paragraphs if p.text == "Full transcript")
        self.assertTrue(transcript.paragraph_format.page_break_before)
        self.assertEqual("Title", document.paragraphs[0].style.name)
        with ZipFile(BytesIO(data)) as archive:
            self.assertIsNone(archive.testzip())
            xml = archive.read("word/document.xml").decode()
            self.assertNotIn("<literal>", xml)
            self.assertIn("&lt;literal&gt;", xml)

    def test_empty_job_does_not_claim_notes_were_generated(self):
        output = render_markdown({})
        self.assertIn("Meeting notes have not been generated", output)
        self.assertIn("No transcript is available", output)
        self.assertNotIn("None identified", output)
        Document(BytesIO(render_docx({})))

    def test_consecutive_sentences_form_speaker_paragraphs_in_both_exports(self):
        job = deepcopy(FIXTURE)
        job["speakers"] = [{"id": "A", "name": "Same name"},
                           {"id": "B", "name": "Same name"}]
        job["segments"] = [
            {"id": index, "start": index * 5, "end": index * 5 + 4,
             "speaker": speaker, "text": text}
            for index, (speaker, text) in enumerate([
                ("A", "First sentence."), ("A", "Second sentence."),
                ("B", "Third sentence."), ("B", "Fourth sentence."),
                ("A", "Fifth sentence."),
            ])
        ]
        original = deepcopy(job)
        markdown = render_markdown(job)
        transcript = markdown.split("## Full transcript", 1)[1]
        self.assertIn("First sentence\\. Second sentence\\.", transcript)
        self.assertIn("Third sentence\\. Fourth sentence\\.", transcript)
        self.assertEqual(3, transcript.count("— Same name**"))
        self.assertNotIn("**00:05", transcript)
        self.assertIn("**00:10", transcript)
        self.assertIn("**00:20", transcript)
        # Notes still cite original sentence timestamps, even within a paragraph.
        self.assertIn("Source: 00:05", markdown)

        document = Document(BytesIO(render_docx(job)))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        start = paragraphs.index("Timestamps refer to the original recording.") + 1
        self.assertEqual([
            "00:00  Same name\nFirst sentence. Second sentence.",
            "00:10  Same name\nThird sentence. Fourth sentence.",
            "00:20  Same name\nFifth sentence.",
        ], paragraphs[start:])
        self.assertEqual(original, job)

    def test_paragraphs_break_on_pause_or_length_without_splitting_source(self):
        segments = [
            {"id": 0, "start": 0, "end": 5, "speaker": "A", "text": "a" * 445},
            {"id": 1, "start": 5, "end": 10, "speaker": "A", "text": "b" * 454},
            {"id": 2, "start": 10, "end": 15, "speaker": "A", "text": "c"},
            {"id": 3, "start": 23, "end": 30, "speaker": "A", "text": "After an eight second pause."},
            {"id": 4, "start": 37.9, "end": 40, "speaker": "A", "text": "Shorter pause stays together."},
            {"id": 5, "start": 40, "end": 45, "speaker": "A", "text": "x" * 1000},
        ]
        paragraphs = _transcript_paragraphs(segments)
        self.assertEqual([0, 10, 23, 40], [p["start"] for p in paragraphs])
        self.assertEqual(900, len(paragraphs[0]["text"]))
        self.assertEqual("After an eight second pause. Shorter pause stays together.", paragraphs[2]["text"])
        self.assertEqual("x" * 1000, paragraphs[3]["text"])

    def test_edited_slices_preserve_word_boundaries_line_breaks_and_empty_parts(self):
        job = deepcopy(FIXTURE)
        job["segments"] = [
            {"id": i, "start": i * 2, "end": i * 2 + 2, "speaker": "SPEAKER_00",
             "text": text, "text_join_before": "", "runs": [{"text": text}] if text else []}
            for i, text in enumerate(["  Hel", "lo ", "", "world\n\nNext line.  "])
        ]
        original = deepcopy(job)
        paragraph = _transcript_paragraphs(job["segments"])[0]
        self.assertEqual("  Hello world\n\nNext line.  ", paragraph["text"])
        self.assertEqual(paragraph["text"], "".join(run["text"] for run in paragraph["runs"]))
        markdown = render_markdown(job).split("## Full transcript", 1)[1]
        self.assertIn("  Hello world  \n  \nNext line\\.  ", markdown)
        self.assertNotIn("[No text]", markdown)
        document = Document(BytesIO(render_docx(job)))
        self.assertEqual("00:00  Example Alex\n  Hello world\n\nNext line.  ", document.paragraphs[-1].text)
        self.assertEqual(job, original)

    def test_bold_italic_underline_export_without_interpreting_source_markup(self):
        job = deepcopy(FIXTURE)
        runs = [{"text": "Bold", "bold": True}, {"text": " & "},
                {"text": "italic", "italic": True}, {"text": "\n"},
                {"text": "<script>literal</script>", "bold": True, "italic": True, "underline": True}]
        job["segments"] = [{"id": 0, "start": 0, "end": 15, "speaker": "SPEAKER_00",
                            "text": "".join(run["text"] for run in runs),
                            "text_join_before": "", "runs": runs}]
        markdown = render_markdown(job)
        self.assertIn("<strong>Bold</strong>", markdown)
        self.assertIn("<em>italic</em>", markdown)
        self.assertIn("<strong><em><u>&lt;script&gt;literal&lt;/script&gt;</u></em></strong>", markdown)
        self.assertNotIn("<script>", markdown)
        document = Document(BytesIO(render_docx(job)))
        body_runs = document.paragraphs[-1].runs[2:]
        self.assertEqual("".join(run["text"] for run in runs), "".join(run.text for run in body_runs))
        self.assertTrue(next(run for run in body_runs if run.text == "Bold").bold)
        self.assertTrue(next(run for run in body_runs if run.text == "italic").italic)
        literal = next(run for run in body_runs if run.text.startswith("<script>"))
        self.assertTrue(literal.bold and literal.italic and literal.underline)

    def test_explicit_prefix_is_ignored_at_paragraph_start_and_legacy_runs_trim(self):
        segments = [
            {"id": 0, "start": 0, "end": 1, "speaker": "A", "text": "  Word  ",
             "runs": [{"text": "  Word  ", "bold": True}]},
            {"id": 1, "start": 1, "end": 2, "speaker": "A", "text": "joined", "text_join_before": " "},
            {"id": 2, "start": 2, "end": 3, "speaker": "B", "text": "New speaker", "text_join_before": " "},
        ]
        paragraphs = _transcript_paragraphs(segments)
        self.assertEqual(["Word joined", "New speaker"], [paragraph["text"] for paragraph in paragraphs])
        self.assertEqual("Word", paragraphs[0]["runs"][0]["text"])
        self.assertTrue(paragraphs[0]["runs"][0]["bold"])

    def test_invalid_stored_runs_fall_back_to_complete_plain_text(self):
        segment = {"id": 0, "speaker": "A", "start": 0, "end": 1,
                   "text": "Keep all these words", "text_join_before": "",
                   "runs": [{"text": "Only part", "bold": True}]}
        paragraph = _transcript_paragraphs([segment])[0]
        self.assertEqual("Keep all these words", "".join(run["text"] for run in paragraph["runs"]))
        self.assertFalse(paragraph["runs"][0]["bold"])

    def test_long_edited_passage_keeps_raw_fragments_and_empty_slices_together(self):
        job = deepcopy(FIXTURE)
        fragments = ["A long edited passage. " * 60 + "inter", "", "national\nFinal line."]
        job["segments"] = [
            {"id": index, "start": index * 5, "end": index * 5 + 5,
             "speaker": "SPEAKER_00", "text": text, "text_join_before": "",
             "runs": [{"text": text, "bold": True}] if text else []}
            for index, text in enumerate(fragments)
        ]
        expected = "".join(fragments)
        paragraphs = _transcript_paragraphs(job["segments"])
        self.assertEqual(1, len(paragraphs))
        self.assertEqual(expected, paragraphs[0]["text"])
        self.assertIn("international", paragraphs[0]["text"])
        markdown = render_markdown(job).split("## Full transcript", 1)[1]
        self.assertEqual(1, markdown.count("— Example Alex**"))
        self.assertNotIn("[No text]", markdown)
        document = Document(BytesIO(render_docx(job)))
        self.assertEqual("00:00  Example Alex\n" + expected, document.paragraphs[-1].text)

        # Explicit spaces, speaker changes and long pauses still create boundaries.
        for changes in ({"text_join_before": " "}, {"speaker": "SPEAKER_01"}, {"start": 13}):
            changed = deepcopy(job["segments"])
            changed[1].update(changes)
            self.assertGreater(len(_transcript_paragraphs(changed)), 1)


if __name__ == "__main__":
    if "--sample" in sys.argv:
        sample = Path(__file__).with_name("synthetic-example.docx")
        sample.write_bytes(render_docx(FIXTURE))
        sample.with_suffix(".md").write_text(render_markdown(FIXTURE), encoding="utf-8")
        print(f"Saved explicitly synthetic QA example to {sample}")
    else:
        unittest.main()

