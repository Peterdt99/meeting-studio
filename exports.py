"""Readable, self-contained meeting exports. No model or network access occurs here."""
from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
import math
import re
from typing import Any
from languages import LANGUAGE_NAMES


_DRAFT_NOTICE = (
    "AI draft. Speaker labels are estimated and names are assigned manually. "
    "Check the transcript, speaker attribution and meeting notes before sharing."
)
_STALE_NOTICE = (
    "Notes need updating. The transcript or speaker names changed after these notes "
    "were generated. Regenerate the notes before sharing."
)


def _text(value: Any) -> str:
    """Keep literal user content while dropping characters illegal in XML 1.0."""
    if value is None:
        return ""
    return "".join(
        ch for ch in str(value)
        if ch in "\t\n\r" or "\x20" <= ch <= "\ud7ff"
        or "\ue000" <= ch <= "\ufffd" or "\U00010000" <= ch <= "\U0010ffff"
    ).replace("\r\n", "\n").replace("\r", "\n")


def _md(value: Any) -> str:
    """Escape untrusted text rather than interpreting it as Markdown or HTML."""
    value = _text(value).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return re.sub(r"([\\`*_{}\[\]()#+.!|\-])", r"\\\1", value).replace("\n", "  \n")


def _seconds(value: Any) -> float | None:
    try:
        number = float(value)
        return max(0.0, number) if math.isfinite(number) else None
    except (TypeError, ValueError, OverflowError):
        return None


def _time(value: Any) -> str:
    seconds = _seconds(value)
    if seconds is None:
        return "Time unavailable"
    total = int(seconds)
    hours, remaining = divmod(total, 3600)
    minutes, seconds = divmod(remaining, 60)
    return f"{hours:02d}:{minutes:02d}:{seconds:02d}" if hours else f"{minutes:02d}:{seconds:02d}"


def _date(value: Any) -> str:
    raw = _text(value).strip()
    if not raw:
        return "Not provided"
    try:
        parsed = datetime.fromisoformat(raw.replace("Z", "+00:00"))
    except ValueError:
        return raw
    # The job timestamp is the import time, never an inferred recording date.
    result = parsed.strftime("%B %d, %Y at %H:%M")
    if parsed.utcoffset() is not None:
        offset = parsed.strftime("%z")
        result += " UTC" if offset == "+0000" else f" UTC{offset[:3]}:{offset[3:]}"
    return result


def _language(value: Any) -> str:
    raw = _text(value).strip()
    return LANGUAGE_NAMES.get(raw.lower(), raw or "Not provided")


def _speakers(job: dict) -> dict[str, str]:
    return {
        _text(speaker.get("id")): _text(speaker.get("name")).strip()
        or _text(speaker.get("id")).strip() or "Unassigned speaker"
        for speaker in job.get("speakers", []) if isinstance(speaker, dict)
    }


def _speaker(value: Any, names: dict[str, str]) -> str:
    key = _text(value)
    return names.get(key, key or "Unassigned speaker")


def _source(item: dict, segments: dict[str, dict]) -> str:
    references = []
    missing = False
    for segment_id in item.get("segment_ids", []) or []:
        key = _text(segment_id)
        if key in references:
            continue
        if key in segments:
            references.append(key)
        else:
            missing = True
    times = [_time(segments[key].get("start")) for key in references]
    label = "Source: " + ", ".join(times) if times else "Source reference unavailable"
    return label + ("; some references unavailable" if missing and times else "")


def _transcript_paragraphs(segments: list[dict]) -> list[dict]:
    """Group turns by speaker, pause and length without changing source segments."""
    paragraphs = []
    for segment in segments:
        speaker = _text(segment.get("speaker"))
        text, runs = _segment_display(segment)
        separator = segment.get("text_join_before")
        if separator not in ("", " "):
            separator = " "
        previous = paragraphs[-1] if paragraphs else None
        start = _seconds(segment.get("start"))
        previous_end = _seconds(previous["end"]) if previous else None
        long_pause = (start is not None and previous_end is not None
                      and start - previous_end >= 8)
        # Empty explicit separators stitch slices of one edited passage, including
        # empty slices retained for original timestamps. Never split these mid-word.
        stitched_slice = segment.get("text_join_before") == ""
        if (previous and previous["speaker"] == speaker and not long_pause
                and (stitched_slice or previous["length"] + len(separator) + len(text) <= 900)):
            previous["text"] += separator + text
            _append_run(previous["runs"], {"text": separator})
            for run in runs:
                _append_run(previous["runs"], run)
            previous["length"] += len(separator) + len(text)
            previous["end"] = segment.get("end")
        else:
            paragraphs.append({"start": segment.get("start"), "end": segment.get("end"),
                               "speaker": speaker, "text": text, "runs": runs,
                               "length": len(text)})
    return [{"start": paragraph["start"], "speaker": paragraph["speaker"],
             "text": paragraph["text"], "runs": paragraph["runs"]}
            for paragraph in paragraphs]


def _append_run(target: list[dict], run: dict) -> None:
    text = _text(run.get("text"))
    if not text:
        return
    flags = {flag: run.get(flag) is True for flag in ("bold", "italic", "underline")}
    if target and all(target[-1][flag] == value for flag, value in flags.items()):
        target[-1]["text"] += text
    else:
        target.append({"text": text, **flags})


def _segment_display(segment: dict) -> tuple[str, list[dict]]:
    """Keep edited slices exact; legacy segments retain their original trimming."""
    text = _text(segment.get("text"))
    source_runs = segment.get("runs")
    if (not isinstance(source_runs, list)
            or any(not isinstance(run, dict) for run in source_runs)
            or "".join(_text(run.get("text")) for run in source_runs) != text):
        source_runs = [{"text": text}]
    left, right = 0, len(text)
    if segment.get("text_join_before") not in ("", " "):
        left, right = len(text) - len(text.lstrip()), len(text.rstrip())
    runs, position = [], 0
    for source in source_runs:
        run_text = _text(source.get("text"))
        start, end = max(left - position, 0), min(right - position, len(run_text))
        if end > start:
            _append_run(runs, {**source, "text": run_text[start:end]})
        position += len(run_text)
    return text[left:right], runs


def _markdown_runs(runs: list[dict]) -> str:
    """Only generated inline style tags are HTML; all source content is escaped."""
    output = []
    for run in runs:
        # Close tags at line boundaries so blank lines never cross HTML elements.
        for part in re.split(r"(\n+)", run["text"]):
            if not part:
                continue
            escaped = _md(part)
            if not part.startswith("\n"):
                for flag, tag in (("underline", "u"), ("italic", "em"), ("bold", "strong")):
                    if run.get(flag):
                        escaped = f"<{tag}>{escaped}</{tag}>"
            output.append(escaped)
    return "".join(output)


def _view(job: dict) -> dict:
    names = _speakers(job)
    segments = [s for s in job.get("segments", []) if isinstance(s, dict)]
    by_id = {_text(s.get("id")): s for s in segments}
    used_speakers = list(dict.fromkeys(_text(s.get("speaker")) for s in segments))
    listed_names = list(dict.fromkeys(names.get(key, key) for key in used_speakers
                                     if key != "UNKNOWN" or names.get(key) not in {None, "Unassigned"}))
    warnings = [_text(w).strip() for w in job.get("warnings", []) if _text(w).strip()]
    if "UNKNOWN" not in used_speakers:
        warnings = [w for w in warnings if not w.startswith("Some words could not be confidently matched")]
    notes = job.get("notes") if isinstance(job.get("notes"), dict) else None
    return {
        "title": _text(job.get("title")).strip() or "Meeting record",
        "metadata": [
            ("Imported", _date(job.get("created_at"))),
            ("Recording", _text(job.get("filename")).strip() or "Not provided"),
            ("Duration", _time(job.get("duration"))),
            ("Language", _language(job.get("language"))),
            ("Speakers", ", ".join(listed_names) or "Not assigned"),
        ],
        "names": names, "segments": segments, "by_id": by_id, "notes": notes,
        "warnings": warnings,
    }


def render_markdown(job: dict) -> str:
    """Return a complete meeting document with literal escaped source content."""
    view = _view(job)
    lines = ["# " + _md(view["title"]), ""]
    for label, value in view["metadata"]:
        lines.extend([f"**{label}:** {_md(value)}", ""])
    lines.extend([_DRAFT_NOTICE, ""])
    if job.get("notes_stale"):
        lines.extend(["**" + _STALE_NOTICE + "**", ""])
    if view["warnings"]:
        lines.extend(["## Processing notes", ""])
        lines.extend(["- " + _md(warning) for warning in view["warnings"]])
        lines.append("")
    lines.extend(["## Meeting notes", ""])
    notes = view["notes"]
    if notes is None:
        lines.extend(["Meeting notes have not been generated.", ""])
    else:
        lines.extend([_md(notes.get("overview")) or "No overview was generated.", ""])
        for key, label in [("decisions", "Decisions"), ("actions", "Action items"),
                           ("open_questions", "Open questions")]:
            lines.extend(["### " + label, ""])
            items = [item for item in notes.get(key, []) or [] if isinstance(item, dict)]
            if not items:
                lines.extend(["None identified in the generated notes.", ""])
            for item in items:
                lines.append("- " + (_md(item.get("text")) or "No text provided."))
                if key == "actions":
                    owner = _speaker(item.get("owner"), view["names"]) if item.get("owner") else "Not specified"
                    lines.append("  - Owner: " + _md(owner) + "; due: " + _md(item.get("due") or "Not specified"))
                lines.extend(["  - " + _md(_source(item, view["by_id"])), ""])
    lines.extend(["## Full transcript", "", "Timestamps refer to the original recording.", ""])
    if not view["segments"]:
        lines.extend(["No transcript is available.", ""])
    for segment in _transcript_paragraphs(view["segments"]):
        lines.extend([
            "**" + _md(_time(segment.get("start"))) + " — "
            + _md(_speaker(segment.get("speaker"), view["names"])) + "**",
            "", _markdown_runs(segment["runs"]), "",
        ])
    # Keep intentional trailing spaces in an edited final passage.
    return "\n".join(lines) + "\n"


def render_docx(job: dict) -> bytes:
    """Return an editable DOCX meeting report with a separate transcript page."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    view = _view(job)
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.85)
    section.footer_distance = Inches(0.3)

    for style_name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "List Bullet"]:
        style = document.styles[style_name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.12
        style.paragraph_format.widow_control = True
    document.styles["Title"].font.size = Pt(25)
    document.styles["Title"].font.bold = True
    document.styles["Title"].paragraph_format.space_after = Pt(14)
    for name, size in [("Heading 1", 15), ("Heading 2", 12)]:
        style = document.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(13)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    core = document.core_properties
    core.title = view["title"]
    core.subject = "Meeting notes and transcript"
    core.author = "Meeting Studio"
    core.last_modified_by = "Meeting Studio"
    core.comments = ""
    core.keywords = ""
    core.created = core.modified = datetime.now(timezone.utc)

    document.add_paragraph(view["title"], style="Title")
    for label, value in view["metadata"]:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.add_run(label + ": ").bold = True
        paragraph.add_run(value)

    paragraph = document.add_paragraph(_DRAFT_NOTICE)
    paragraph.paragraph_format.space_before = Pt(10)
    for run in paragraph.runs:
        run.italic = True
        run.font.size = Pt(9)
    if job.get("notes_stale"):
        paragraph = document.add_paragraph()
        paragraph.add_run(_STALE_NOTICE).bold = True
    if view["warnings"]:
        document.add_heading("Processing notes", 1)
        for warning in view["warnings"]:
            document.add_paragraph(warning, style="List Bullet")

    document.add_heading("Meeting notes", 1)
    notes = view["notes"]
    if notes is None:
        document.add_paragraph("Meeting notes have not been generated.")
    else:
        overview = _text(notes.get("overview")).strip() or "No overview was generated."
        for text in overview.split("\n\n"):
            document.add_paragraph(text)
        for key, label in [("decisions", "Decisions"), ("actions", "Action items"),
                           ("open_questions", "Open questions")]:
            document.add_heading(label, 2)
            items = [item for item in notes.get(key, []) or [] if isinstance(item, dict)]
            if not items:
                document.add_paragraph("None identified in the generated notes.")
            for item in items:
                paragraph = document.add_paragraph(_text(item.get("text")) or "No text provided.", style="List Bullet")
                paragraph.paragraph_format.keep_with_next = True
                if key == "actions":
                    owner = _speaker(item.get("owner"), view["names"]) if item.get("owner") else "Not specified"
                    paragraph = document.add_paragraph()
                    paragraph.paragraph_format.left_indent = Inches(0.25)
                    paragraph.paragraph_format.space_after = Pt(3)
                    paragraph.paragraph_format.keep_with_next = True
                    paragraph.add_run("Owner: ").bold = True
                    paragraph.add_run(owner + "    ")
                    paragraph.add_run("Due: ").bold = True
                    paragraph.add_run(_text(item.get("due")) or "Not specified")
                paragraph = document.add_paragraph(_source(item, view["by_id"]))
                paragraph.paragraph_format.left_indent = Inches(0.25)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(85, 85, 85)

    heading = document.add_heading("Full transcript", 1)
    heading.paragraph_format.page_break_before = True
    document.add_paragraph("Timestamps refer to the original recording.")
    if not view["segments"]:
        document.add_paragraph("No transcript is available.")
    for segment in _transcript_paragraphs(view["segments"]):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(10)
        paragraph.paragraph_format.keep_together = False
        paragraph.add_run(_time(segment.get("start")) + "  " + _speaker(segment.get("speaker"), view["names"])).bold = True
        paragraph.add_run("\n")
        for source_run in segment["runs"]:
            run = paragraph.add_run(source_run["text"])
            run.bold = source_run["bold"]
            run.italic = source_run["italic"]
            run.underline = source_run["underline"]

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Page ")
    run.font.size = Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    output = BytesIO()
    document.save(output)
    return output.getvalue()
